"""XLSX exporter for alignment data and validation reports."""
from __future__ import annotations

import math
from pathlib import Path

from railtrack.domain.alignment import Alignment
from railtrack.domain.horizontal import CircularArc, Straight, TransitionCurve
from railtrack.domain.validation_types import ValidationIssue


def export_alignment_xlsx(
    alignment: Alignment,
    path: Path | str,
    issues: list[ValidationIssue] | None = None,
) -> None:
    """Export alignment data and optionally validation issues to XLSX."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()

    # --- Horizontal elements sheet ---
    ws_h = wb.active
    ws_h.title = "Horizontal"
    headers_h = [
        "Index", "Type", "Start Ch", "End Ch", "Length",
        "Start X", "Start Y", "End X", "End Y",
        "Start Az (°)", "End Az (°)",
        "Radius", "Direction",
    ]
    ws_h.append(headers_h)
    for cell in ws_h[1]:
        cell.font = Font(bold=True)

    for i, elem in enumerate(alignment.horizontal_elements):
        etype = "straight"
        radius = ""
        direction = ""
        if isinstance(elem, CircularArc):
            etype = "arc"
            radius = round(elem.radius, 2)
            direction = elem.direction.value
        elif isinstance(elem, TransitionCurve):
            etype = "transition"
            direction = elem.direction.value

        ws_h.append([
            i, etype,
            round(elem.start_chainage, 4), round(elem.end_chainage, 4),
            round(elem.length, 4),
            round(elem.start_point.x, 4), round(elem.start_point.y, 4),
            round(elem.end_point.x, 4), round(elem.end_point.y, 4),
            round(math.degrees(elem.start_azimuth), 6),
            round(math.degrees(elem.end_azimuth), 6),
            radius, direction,
        ])

    # --- Vertical elements sheet ---
    if alignment.vertical_elements:
        ws_v = wb.create_sheet("Vertical")
        ws_v.append(["Index", "Type", "Start Ch", "End Ch", "Length", "Start Elev", "End Elev", "Gradient (‰)"])
        for cell in ws_v[1]:
            cell.font = Font(bold=True)
        for i, elem in enumerate(alignment.vertical_elements):
            ws_v.append([
                i, elem.element_type.value,
                round(elem.start_chainage, 4),
                round(elem.end_chainage, 4),
                round(elem.length, 4),
                round(elem.elevation_at(elem.start_chainage), 4),
                round(elem.elevation_at(elem.end_chainage), 4),
                round(elem.gradient_at(elem.start_chainage) * 1000, 2),
            ])

    # --- Validation issues sheet ---
    if issues:
        ws_val = wb.create_sheet("Validation")
        ws_val.append(["Severity", "Code", "Title", "Message", "Chainage", "Expected", "Actual", "Suggestion"])
        for cell in ws_val[1]:
            cell.font = Font(bold=True)

        severity_fills = {
            "error": PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid"),
            "warning": PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid"),
            "info": PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid"),
        }

        for issue in issues:
            row_idx = ws_val.max_row + 1
            ws_val.append([
                issue.severity.value,
                issue.code,
                issue.title,
                issue.message,
                issue.chainage_display,
                issue.expected_value,
                issue.actual_value,
                issue.suggestion,
            ])
            fill = severity_fills.get(issue.severity.value)
            if fill:
                for cell in ws_val[row_idx]:
                    cell.fill = fill

    wb.save(path)
